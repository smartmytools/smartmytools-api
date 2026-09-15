import io
import subprocess
import tempfile
from pathlib import Path

import fitz

from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_SECTION


# =========================================
# LIBREOFFICE PATH
# =========================================

SOFFICE_PATH = Path(
    r"C:\Program Files\LibreOffice\program\soffice.exe"
)


# =========================================
# MAIN CONVERSION
# =========================================

def convert_pdf_to_word(pdf_bytes):

    if not pdf_bytes:
        raise ValueError("The PDF file is empty.")

    # -------------------------------------
    # Try LibreOffice first
    # -------------------------------------

    try:

        result = convert_with_libreoffice(
            pdf_bytes
        )

        if result:
            return result

    except Exception as error:

        print(
            "LibreOffice conversion failed:",
            repr(error)
        )

    # -------------------------------------
    # Fallback:
    # Preserve each PDF page as a Word page
    # -------------------------------------

    print(
        "Using PDF page preservation fallback..."
    )

    return convert_pdf_pages_to_docx(
        pdf_bytes
    )


# =========================================
# LIBREOFFICE CONVERSION
# =========================================

def convert_with_libreoffice(pdf_bytes):

    if not SOFFICE_PATH.exists():

        raise FileNotFoundError(
            "LibreOffice was not found."
        )

    with tempfile.TemporaryDirectory() as temp_dir:

        temp_path = Path(temp_dir)

        input_pdf = (
            temp_path / "document.pdf"
        )

        input_pdf.write_bytes(
            pdf_bytes
        )

        output_dir = (
            temp_path / "output"
        )

        output_dir.mkdir()

        command = [

            str(SOFFICE_PATH),

            "--headless",

            "--convert-to",

            "docx",

            "--outdir",

            str(output_dir),

            str(input_pdf)

        ]

        process = subprocess.run(

            command,

            capture_output=True,

            text=True,

            timeout=120

        )

        print(
            "LibreOffice stdout:",
            process.stdout
        )

        print(
            "LibreOffice stderr:",
            process.stderr
        )

        output_docx = (
            output_dir / "document.docx"
        )

        if output_docx.exists():

            return output_docx.read_bytes()

        return None


# =========================================
# PDF PAGE PRESERVATION FALLBACK
# =========================================

def convert_pdf_pages_to_docx(pdf_bytes):

    pdf_document = fitz.open(

        stream=pdf_bytes,

        filetype="pdf"

    )

    document = Document()

    first_section = (
        document.sections[0]
    )

    for page_number, page in enumerate(
        pdf_document
    ):

        page_width = (
            page.rect.width
        )

        page_height = (
            page.rect.height
        )

        # ---------------------------------
        # Create matching page
        # ---------------------------------

        if page_number == 0:

            section = (
                first_section
            )

        else:

            section = (
                document.add_section(
                    WD_SECTION.NEW_PAGE
                )
            )

        # ---------------------------------
        # Match PDF page size
        # ---------------------------------

        section.page_width = Inches(
            page_width / 72
        )

        section.page_height = Inches(
            page_height / 72
        )

        section.top_margin = Inches(0)

        section.bottom_margin = Inches(0)

        section.left_margin = Inches(0)

        section.right_margin = Inches(0)

        # ---------------------------------
        # Render PDF page
        # ---------------------------------

        pix = page.get_pixmap(

            matrix=fitz.Matrix(2, 2),

            alpha=False

        )

        image_stream = io.BytesIO(

            pix.tobytes("png")

        )

        # ---------------------------------
        # Add page image
        # ---------------------------------

        paragraph = (
            document.add_paragraph()
        )

        paragraph.paragraph_format.space_before = 0

        paragraph.paragraph_format.space_after = 0

        run = paragraph.add_run()

        run.add_picture(

            image_stream,

            width=Inches(
                page_width / 72
            ),

            height=Inches(
                page_height / 72
            )

        )

    pdf_document.close()

    # -------------------------------------
    # Save DOCX
    # -------------------------------------

    final_stream = io.BytesIO()

    document.save(
        final_stream
    )

    final_stream.seek(0)

    return final_stream.getvalue()