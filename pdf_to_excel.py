import io
import re

import pymupdf
import pdfplumber
import openpyxl

from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def convert_number(value):
    if not isinstance(value, str):
        return value

    value = value.strip()

    if not value:
        return None

    # Remove superscript footnote markers
    value = re.sub(r"[¹²³⁴⁵⁶⁷⁸⁹]+$", "", value)

    # Remove footnote markers such as 44.7(3)
    value = re.sub(r"\(\d+\)$", "", value)

    value = value.strip()

    cleaned_value = value.replace(",", "")

    if re.fullmatch(r"-?\d+", cleaned_value):
        return int(cleaned_value)

    if re.fullmatch(r"-?\d+\.\d+", cleaned_value):
        return float(cleaned_value)

    return value


def clean_text(value):
    if not value:
        return ""

    value = str(value)
    value = value.replace("\n", " ")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def rectangles_overlap(text_bbox, table_bbox):
    text_x0, text_top, text_x1, text_bottom = text_bbox
    table_x0, table_top, table_x1, table_bottom = table_bbox

    return (
        text_x0 < table_x1
        and text_x1 > table_x0
        and text_top < table_bottom
        and text_bottom > table_top
    )


def extract_text_items(page, table_bboxes):
    """
    Extract text blocks from PyMuPDF while ignoring text
    that is inside detected tables.
    """

    text_items = []

    for block in page.get_text("blocks"):
        if len(block) < 5:
            continue

        x0, top, x1, bottom, text = block[:5]

        text = clean_text(text)

        if not text:
            continue

        text_bbox = (
            x0,
            top,
            x1,
            bottom
        )

        inside_table = False

        for table_bbox in table_bboxes:
            if rectangles_overlap(
                text_bbox,
                table_bbox
            ):
                inside_table = True
                break

        if inside_table:
            continue

        text_items.append({
            "type": "text",
            "top": top,
            "bottom": bottom,
            "text": text
        })

    return text_items


def write_text_item(
    worksheet,
    current_row,
    text,
    is_first_text=False
):
        # Center each table heading over its own table

    if text.startswith(("Table 3:", "Table 4:")):
        heading_start_column = 2
        heading_end_column = 3

    elif text.startswith("Table 5:"):
        heading_start_column = 1
        heading_end_column = 4

    elif text.startswith("Table 6:"):
        heading_start_column = 1
        heading_end_column = 5

    else:
        heading_start_column = None
        heading_end_column = None

    if heading_start_column is not None:
        cell = worksheet.cell(
            row=current_row,
            column=heading_start_column,
            value=text
        )

        worksheet.merge_cells(
            start_row=current_row,
            start_column=heading_start_column,
            end_row=current_row,
            end_column=heading_end_column
        )

        cell.font = Font(
            bold=True,
            size=12
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        worksheet.row_dimensions[
            current_row
        ].height = 24

        return current_row + 1
        cell = worksheet.cell(
            row=current_row,
            column=2,
            value=text
        )

        worksheet.merge_cells(
            start_row=current_row,
            start_column=2,
            end_row=current_row,
            end_column=6
        )

        cell.font = Font(
            bold=True,
            size=12
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        worksheet.row_dimensions[
            current_row
        ].height = 24

        return current_row + 1

    cell = worksheet.cell(
        row=current_row,
        column=1,
        value=text
    )

    if is_first_text:
        cell.font = Font(
            bold=True,
            size=14
        )

        worksheet.row_dimensions[
            current_row
        ].height = 30

    else:
        cell.font = Font(
            bold=True,
            size=12
        )

        line_count = max(
            1,
            (len(text) // 45) + 1
        )

        worksheet.row_dimensions[
            current_row
        ].height = 20 * line_count

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True
    )

    return current_row + 1


def write_table(
    worksheet,
    table,
    current_row,
    start_column=1
):
    thin_side = Side(
        style="thin",
        color="000000"
    )

    table_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    table_start_row = current_row

    for row in table:
        if not row:
            continue

        cleaned_row = [
            convert_number(cell)
            for cell in row
        ]

        for column_index, cell_value in enumerate(
            cleaned_row
        ):
            column_number = (
                start_column + column_index
            )

            cell = worksheet.cell(
                row=current_row,
                column=column_number,
                value=cell_value
            )

            cell.border = table_border

            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True
            )

            if current_row == table_start_row:
                cell.font = Font(
                    bold=True
                )

            if isinstance(cell_value, (int, float)):
                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="center"
                )

        current_row += 1

    return current_row


def convert_pdf_to_excel(pdf_bytes):
    if not pdf_bytes:
        raise ValueError("The PDF file is empty.")

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Converted Data"

    extracted_any_table = False
    current_row = 1
    first_text_written = False

    try:
        with (
            pdfplumber.open(io.BytesIO(pdf_bytes)) as plumber_pdf,
            pymupdf.open(stream=pdf_bytes, filetype="pdf") as mupdf_pdf
        ):

            for page_number, plumber_page in enumerate(
                plumber_pdf.pages,
                start=1
            ):
                mupdf_page = mupdf_pdf[
                    page_number - 1
                ]

                found_tables = plumber_page.find_tables()

                if not found_tables:
                    continue

                extracted_any_table = True

                table_bboxes = [
                    table.bbox
                    for table in found_tables
                ]

                text_items = extract_text_items(
                    mupdf_page,
                    table_bboxes
                )

                page_items = []

                for text_item in text_items:
                    page_items.append(text_item)

                for table_index, found_table in enumerate(
                    found_tables
                ):
                    table = found_table.extract()

                    if not table:
                        continue

                    page_items.append({
                        "type": "table",
                        "top": found_table.bbox[1],
                        "bottom": found_table.bbox[3],
                        "table": table,
                        "table_bbox": found_table.bbox,
                        "table_width": (
                            found_table.bbox[2]
                            - found_table.bbox[0]
                        )
                    })

                # Sort headings, footnotes, and tables
                # according to their vertical position
                page_items.sort(
                    key=lambda item: (
                        item["top"],
                        0 if item["type"] == "text" else 1
                    )
                )

                for item in page_items:

                    if item["type"] == "text":
                        text = item["text"]

                        # Ignore isolated page numbers
                        if re.fullmatch(r"\d+", text):
                            continue

                        current_row = write_text_item(
                            worksheet,
                            current_row,
                            text,
                            is_first_text=(
                                not first_text_written
                            )
                        )

                        first_text_written = True

                    elif item["type"] == "table":
                        current_row += 1

                        # Keep wide tables on the left.
                        # Center the small two-column table, such as Table 3.

                        # Position tables according to their number of columns.
                        # 2-column tables are centered.
                        # Wider tables start from column A.

                        table_rows = item["table"]

                        table_column_count = max(
                            len(row)
                            for row in table_rows
                            if row
                        )

                        if table_column_count <= 2:
                            start_column = 2
                        else:
                            start_column = 1

                        current_row = write_table(
                            worksheet,
                            item["table"],
                            current_row,
                            start_column
                        )

                        current_row += 2

        if not extracted_any_table:
            raise ValueError(
                "No tables were found in this PDF. "
                "Please upload a PDF containing selectable tables."
            )

        # Merge the first title across several columns
        if worksheet.max_row >= 1:
            first_cell = worksheet["A1"]

            if first_cell.value:
                worksheet.merge_cells(
                    start_row=1,
                    start_column=1,
                    end_row=1,
                    end_column=4
                )

                first_cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    wrap_text=True
                )

        # Set useful column widths
        for column_number in range(
            1,
            worksheet.max_column + 1
        ):
            max_length = 0

            for row_number in range(
                1,
                worksheet.max_row + 1
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_number
                )

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            column_letter = get_column_letter(
                column_number
            )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 3, 12),
                35
            )

        worksheet.freeze_panes = "A2"

        output_stream = io.BytesIO()

        workbook.save(output_stream)

        output_stream.seek(0)

        return output_stream.getvalue()

    finally:
        workbook.close()